"""
FCC Near-Field Test wizard for QtTinySA.

Guides the operator through: checklist → arm device-side max hold → unplug →
timed hold → replug → dump frequencies/data 2 → E-field plot with ambient overlay.
"""

from __future__ import annotations

import csv
import json
import logging
import re
import time
from datetime import datetime
from pathlib import Path

import numpy as np
import pyqtgraph
import pyqtgraph.exporters  # noqa: F401  — registers ImageExporter
from PySide6 import QtCore, QtWidgets
from PySide6.QtWidgets import QFileDialog, QListWidgetItem, QMessageBox
from serial.tools import list_ports

# TinySA USB IDs (same as Analyser.openPort)
VID = 0x0483
PID = 0x5740

START_FREQ_HZ = 15_000_000
STOP_FREQ_HZ = 1_000_000_000
RBW_KHZ = 30

# Matches the interval Tiny.setCmdQ uses for its command FIFO
FIFO_INTERVAL_MS = 500

FCC_CLASS_B = [
    (30.0, 88.0, 40.0),
    (88.0, 216.0, 43.5),
    (216.0, 960.0, 46.0),
    (960.0, 1000.0, 54.0),
]

STEP_TITLES = [
    "Checklist",
    "Arm Max Hold",
    "Unplug USB",
    "Hold Countdown",
    "Replug USB",
    "Dump Trace",
    "Results",
]

PAGE_CHECKLIST = 0
PAGE_ARM = 1
PAGE_UNPLUG = 2
PAGE_HOLD = 3
PAGE_REPLUG = 4
PAGE_DUMP = 5
PAGE_RESULTS = 6


def sanitize_label(raw: str) -> str:
    safe = "".join(c if c.isalnum() or c in "-_." else "_" for c in raw).strip("_")
    return safe or "unnamed"


def parse_float_lines(text: str) -> list[float]:
    vals = []
    for line in text.replace("-:.0", "-10.0").splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.search(r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?", line)
        if m:
            try:
                vals.append(float(m.group(0)))
            except ValueError:
                pass
    return vals


def _as_emc_result(actual, E, dbuv, scalar: bool):
    if scalar:
        return float(actual), float(E), float(dbuv)
    return actual, E, dbuv


def compute_emc_tekbox(freq_hz, data2_dbm, amp_gain_db: float = 40.0):
    from EMI_Mapper.measurement_chain import characterize_e5_constant_gain

    return characterize_e5_constant_gain(freq_hz, data2_dbm, amp_gain_db)


def compute_emc_schwarzbeck(freq_hz, data2_dbm, ant_factor: float):
    # E (dBuV/m) = dBm + 107 + AF
    data2_dbm = np.asarray(data2_dbm, dtype=float)
    scalar = data2_dbm.ndim == 0
    dbuv = data2_dbm + 107.0 + ant_factor
    E = 10 ** ((dbuv - 120.0) / 20.0) * 1e6  # uV/m
    return _as_emc_result(data2_dbm, E, dbuv, scalar)


def find_tinysa_ports():
    return [p for p in list_ports.comports() if p.vid == VID and p.pid == PID]


class _WizardSerial:
    """
    Owns the analyser under test for the duration of a wizard run.

    QtTinySA 2.x drives up to four analysers and rebuilds every device object
    whenever one connects or disconnects, numbering them by USB port order.
    This wizard deliberately unplugs its analyser mid-run, so it tracks the
    device by serial number: an index-based lookup would silently retarget a
    different analyser after the replug.
    """

    def __init__(self, usb_instr):
        self.usb_instr = usb_instr
        self.sn = None
        self._held = False

    # ------------------------------------------------------------- lookup
    def _connected(self):
        return [d for d in (self.usb_instr.devices or []) if d is not None]

    def device(self):
        if self.sn is None:
            return None
        for dev in self._connected():
            if dev.sn == self.sn:
                return dev
        return None

    def adopt(self):
        """Bind to the analyser under test, preferring an enabled device."""
        devices = self._connected()
        if not devices:
            self.sn = None
            return None
        chosen = next((d for d in devices if getattr(d, "enabled", False)), devices[0])
        self.sn = chosen.sn
        logging.info(f"FCC wizard bound to analyser serial {self.sn} on {chosen.usbPort}")
        return chosen

    def is_open(self):
        dev = self.device()
        return dev is not None and dev.usb is not None and dev.usb.is_open

    # ------------------------------------------------------------ control
    def take(self):
        """
        Stop the sweep and command FIFO so the wizard can own the port.
        Concurrent scanraw and serialWrite on Windows raises
        PermissionError(13) 'The device does not recognize the command.'
        """
        dev = self.device()
        if dev is None:
            return
        self.usb_instr.stop(restart=False)
        try:
            if dev.fifoTimer.isActive():
                dev.fifoTimer.stop()
        except Exception:
            pass
        try:
            if dev.usb and dev.usb.is_open:
                dev.abort()
        except Exception:
            pass
        for _ in range(80):
            if not dev.threadRunning:
                break
            QtCore.QCoreApplication.processEvents()
            time.sleep(0.05)
        try:
            if dev.usb and dev.usb.is_open:
                dev.clearBuffer()
        except Exception:
            pass
        try:
            while dev.fifo.qsize() > 0:
                dev.fifo.get_nowait()
        except Exception:
            pass
        self._held = True

    def release(self):
        if not self._held:
            return
        self.restart_fifo()

    def restart_fifo(self):
        dev = self.device()
        if dev is not None:
            try:
                if not dev.fifoTimer.isActive():
                    dev.fifoTimer.start(FIFO_INTERVAL_MS)
            except Exception:
                pass
        self._held = False

    def close_device(self):
        """Close only the analyser under test, leaving any others running."""
        dev = self.device()
        if dev is not None:
            dev.close()
        self._held = False

    def reprobe(self, timeout_s=10.0):
        """Re-detect USB devices and wait for our serial number to reappear."""
        deadline = time.time() + timeout_s
        while time.time() < deadline:
            try:
                self.usb_instr.probe()
            except Exception as e:
                logging.info(f"FCC reprobe: {e}")
            if self.is_open():
                return self.device()
            QtCore.QCoreApplication.processEvents()
            time.sleep(0.25)
        return self.device()

    # --------------------------------------------------------------- i/o
    def write(self, cmd, retries=2):
        dev = self.device()
        if dev is None:
            raise RuntimeError("Analyser under test is not connected")
        last = None
        for attempt in range(retries + 1):
            try:
                dev.clearBuffer()
                dev.serialWrite(cmd)
                return
            except Exception as e:
                last = e
                logging.info(f"FCC serialWrite retry {attempt}: {cmd!r} -> {e}")
                time.sleep(0.2)
                try:
                    dev.clearBuffer()
                except Exception:
                    pass
        raise last

    def query(self, cmd):
        dev = self.device()
        if dev is None:
            raise RuntimeError("Analyser under test is not connected")
        dev.clearBuffer()
        return dev.serialQuery(cmd)


class FCCWizard(QtCore.QObject):
    """Controller for the FCC Test dialog loaded from fcc_test.ui."""

    def __init__(self, ui, usb_instr, main_window):
        super().__init__()
        self.ui = ui
        self.serial = _WizardSerial(usb_instr)
        self.main = main_window
        self.usbCheck = None  # set by QtTinySA after usbCheck QTimer is created

        self.session_dir: Path | None = None
        self.session_runs: list[dict] = []
        self.armed = False
        self.hold_total_s = 180
        self.hold_remaining_s = 180
        self._quiesced = False
        self._plot_curves: dict[str, object] = {}

        self.poll_timer = QtCore.QTimer(self)
        self.poll_timer.setInterval(500)
        self.poll_timer.timeout.connect(self._on_poll_usb)

        self.countdown_timer = QtCore.QTimer(self)
        self.countdown_timer.setInterval(1000)
        self.countdown_timer.timeout.connect(self._on_countdown_tick)

        self._wire_ui()
        self._setup_plot()

    # ------------------------------------------------------------------ UI
    def _wire_ui(self):
        u = self.ui
        u.btnBack.clicked.connect(self._on_back)
        u.btnNext.clicked.connect(self._on_next)
        u.btnCancel.clicked.connect(self._on_cancel)
        u.btnMeasureSweep.clicked.connect(self._measure_sweep_time)
        u.btnArm.clicked.connect(self._arm_device)
        u.btnAbortHold.clicked.connect(self._abort_hold)
        u.btnNewRun.clicked.connect(self._new_run)
        u.btnExportExcel.clicked.connect(self._export_excel)
        u.btnSavePng.clicked.connect(self._save_png)
        u.runList.itemChanged.connect(self._on_run_visibility)

        for chk in (u.chkPhoto, u.chkDutCharged, u.chkTinyCharged, u.chkPlacement):
            chk.stateChanged.connect(self._update_nav)
        u.runLabel.textChanged.connect(self._update_nav)
        u.ambientRun.stateChanged.connect(self._on_ambient_toggled)
        u.probePath.currentIndexChanged.connect(self._on_probe_path_changed)

        u.wizardStack.currentChanged.connect(self._on_page_changed)

    def _setup_plot(self):
        pw = self.ui.fccPlot
        # Match the Combined EMI reference style: linear MHz x-axis, dBuV/m y-axis
        pw.setBackground("w")
        pw.showGrid(x=True, y=True, alpha=0.35)
        pw.setLogMode(x=False, y=False)
        pw.setLabel("bottom", "Frequency (MHz)")
        pw.setLabel("left", "Electric Field (dBµV/m)")
        for axis_name in ("bottom", "left"):
            axis = pw.getAxis(axis_name)
            axis.enableAutoSIPrefix(False)
            axis.setPen(pyqtgraph.mkPen("k"))
            axis.setTextPen(pyqtgraph.mkPen("k"))
        pw.setXRange(0, 1000, padding=0)
        pw.setYRange(10, 90, padding=0)
        pw.addLegend(offset=(10, 10))
        # Stepped FCC Class B limit (dashed grey), drawn once and kept
        self._fcc_limit_curve = pw.plot(
            pen=pyqtgraph.mkPen((120, 120, 120), width=2.5, style=QtCore.Qt.PenStyle.DashLine),
            name="FCC Class B Limit",
        )
        self._update_fcc_limit_curve()

    def _update_fcc_limit_curve(self):
        # Frequency in MHz (same units as run traces)
        xs, ys = [], []
        for f0, f1, lim in FCC_CLASS_B:
            xs.extend([f0, f1])
            ys.extend([lim, lim])
        # Extend flat to 0 MHz for visual match with reference plots
        if xs:
            xs = [0.0, xs[0]] + xs
            ys = [ys[0], ys[0]] + ys
        self._fcc_limit_curve.setData(xs, ys)

    # ------------------------------------------------------------------ public
    def start(self):
        if self.session_dir is None:
            default = Path.cwd() / f"FCC_Session_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            chosen = QFileDialog.getExistingDirectory(
                self.ui, "Choose / create FCC session folder", str(default.parent)
            )
            if not chosen:
                return
            self.session_dir = Path(chosen)
            self.session_dir.mkdir(parents=True, exist_ok=True)
            self._load_prefs()

        self.armed = False
        self.ui.wizardStack.setCurrentIndex(PAGE_CHECKLIST)
        self._update_header()
        self._update_nav()
        self.ui.show()
        self.ui.raise_()
        self.ui.activateWindow()

    # ------------------------------------------------------------------ prefs
    def _prefs_path(self) -> Path | None:
        if self.session_dir is None:
            return None
        return self.session_dir / "fcc_wizard_prefs.json"

    def _load_prefs(self):
        p = self._prefs_path()
        if not p or not p.is_file():
            return
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
            self.ui.ampGainDb.setValue(float(data.get("amp_gain_db", 40)))
            self.ui.carrierMhz.setValue(float(data.get("carrier_mhz", 20)))
            self.ui.holdMinutes.setValue(int(data.get("hold_minutes", 3)))
            self.ui.probePath.setCurrentIndex(int(data.get("probe_path", 0)))
            self.ui.distanceCm.setValue(float(data.get("distance_cm", 2.0)))
        except Exception as e:
            logging.info(f"FCC prefs load failed: {e}")

    def _save_prefs(self):
        p = self._prefs_path()
        if not p:
            return
        data = {
            "amp_gain_db": self.ui.ampGainDb.value(),
            "carrier_mhz": self.ui.carrierMhz.value(),
            "hold_minutes": self.ui.holdMinutes.value(),
            "probe_path": self.ui.probePath.currentIndex(),
            "distance_cm": self.ui.distanceCm.value(),
        }
        p.write_text(json.dumps(data, indent=2), encoding="utf-8")

    # ------------------------------------------------------------------ nav
    def _update_header(self):
        idx = self.ui.wizardStack.currentIndex()
        self.ui.stepHeader.setText(f"Step {idx + 1} of 7 — {STEP_TITLES[idx]}")

    def _checklist_ok(self) -> bool:
        u = self.ui
        return (
            u.chkPhoto.isChecked()
            and u.chkDutCharged.isChecked()
            and u.chkTinyCharged.isChecked()
            and u.chkPlacement.isChecked()
            and bool(u.runLabel.text().strip())
        )

    def _update_nav(self):
        idx = self.ui.wizardStack.currentIndex()
        self.ui.btnBack.setEnabled(idx in (PAGE_CHECKLIST, PAGE_ARM, PAGE_RESULTS))
        if idx == PAGE_CHECKLIST:
            self.ui.btnNext.setEnabled(self._checklist_ok())
            self.ui.btnNext.setText("Next")
        elif idx == PAGE_ARM:
            self.ui.btnNext.setEnabled(self.armed)
            self.ui.btnNext.setText("Next — Unplug")
        elif idx == PAGE_RESULTS:
            self.ui.btnNext.setEnabled(False)
            self.ui.btnNext.setText("Done")
        else:
            # Auto-advance pages: hide Next during unplug/hold/replug/dump
            self.ui.btnNext.setEnabled(False)
            self.ui.btnNext.setText("Next")

    def _on_page_changed(self, idx: int):
        self._update_header()
        self._update_nav()

    def _on_ambient_toggled(self, state):
        if self.ui.ambientRun.isChecked() and not self.ui.runLabel.text().strip():
            self.ui.runLabel.setText("Ambient")
        self._update_nav()

    def _on_probe_path_changed(self, index: int):
        # TekBox uses amp gain; Schwarzbeck uses AF (amp field becomes AF display hint)
        if index == 0:
            self.ui.ampGainDb.setValue(40.0)
            self.ui.lblAmpGain.setText("Amp gain (dB)")
        elif index == 1:
            self.ui.ampGainDb.setValue(20.0)
            self.ui.lblAmpGain.setText("Antenna factor (dB/m)")
        else:
            self.ui.ampGainDb.setValue(47.0)
            self.ui.lblAmpGain.setText("Antenna factor (dB/m)")

    def _on_back(self):
        idx = self.ui.wizardStack.currentIndex()
        if idx == PAGE_ARM:
            self.ui.wizardStack.setCurrentIndex(PAGE_CHECKLIST)
        elif idx == PAGE_RESULTS:
            # stay; New Run is the way back to checklist
            pass
        elif idx == PAGE_CHECKLIST:
            pass

    def _on_next(self):
        idx = self.ui.wizardStack.currentIndex()
        if idx == PAGE_CHECKLIST:
            if self.ui.ambientRun.isChecked() and not self.ui.runLabel.text().lower().startswith("ambient"):
                # ensure ambient naming for plot ordering
                if "ambient" not in self.ui.runLabel.text().lower():
                    self.ui.runLabel.setText("Ambient_" + self.ui.runLabel.text().strip())
            self.ui.wizardStack.setCurrentIndex(PAGE_ARM)
            self.armed = False
            self.ui.armStatus.setText("Not armed yet.")
            self._update_nav()
        elif idx == PAGE_ARM:
            if not self.armed:
                return
            self._begin_unplug()

    def _on_cancel(self):
        self._cleanup(restore=True)
        self.ui.hide()

    # ------------------------------------------------------------------ arm
    def _require_usb(self) -> bool:
        if self.serial.device() is None:
            self.serial.adopt()
        if not self.serial.is_open():
            QMessageBox.critical(self.ui, "TinySA", "TinySA not connected. Connect the device first.")
            return False
        return True

    def _clear_maxhold_buffer(self):
        """Turn off device calc mode so the previous max-hold trace cannot bleed into the next run."""
        if not self.serial.is_open():
            return
        try:
            self.serial.take()
            self.serial.write("pause\r")
            self.serial.write("calc off\r")
            time.sleep(0.1)
            logging.info("FCC wizard: cleared max-hold (calc off)")
        except Exception as e:
            logging.info(f"FCC wizard: calc off failed: {e}")
        finally:
            # Keep FIFO stopped only if we're mid-wizard; otherwise release
            if self.ui.wizardStack.currentIndex() in (PAGE_CHECKLIST, PAGE_ARM, PAGE_RESULTS):
                self.serial.release()

    def _send_arm_commands(self):
        # Console command is "ultra on" (no PIN). PIN 4321 is only for the
        # on-device Ultra Mode unlock menu; serial does not accept it.
        # RBW 30 kHz (finer than 100 kHz for harmonic visibility).
        cmds = [
            "pause\r",
            "ultra on\r",
            "calc off\r",
            f"sweep start {START_FREQ_HZ}\r",
            f"sweep stop {STOP_FREQ_HZ}\r",
            f"rbw {RBW_KHZ}\r",
            "attenuate 0\r",
            "lna off\r",
            "spur auto\r",
            "calc maxh\r",
            "resume\r",
        ]
        for c in cmds:
            self.serial.write(c)
            time.sleep(0.08)

    def _verify_maxhold(self) -> bool:
        time.sleep(1.0)
        r1 = parse_float_lines(self.serial.query("data 2\r"))
        time.sleep(2.0)
        r2 = parse_float_lines(self.serial.query("data 2\r"))
        if r1 and r2 and len(r1) == len(r2):
            diffs = np.array(r2) - np.array(r1)
            return bool(np.nanmin(diffs) >= -0.5)
        return False

    def _apply_arm_result(self, ok: bool):
        self.armed = True
        if ok:
            self.ui.armStatus.setText(
                "Armed: Ultra on, buffer cleared, fresh Max Hold confirmed. "
                "Press Next to unplug USB."
            )
            self.ui.armStatus.setStyleSheet("color: #0a0;")
            return
        self.ui.armStatus.setText(
            "WARNING: could not verify calc maxh via serial. "
            "On the TinySA screen, enable Max Hold manually (re-select Max Hold to reset), "
            "confirm peaks accumulate, then press Next."
        )
        self.ui.armStatus.setStyleSheet("color: orange;")
        QMessageBox.warning(
            self.ui,
            "Verify Max Hold",
            "Serial max-hold verification inconclusive.\n\n"
            "On the TinySA: Display/Calc → Max Hold (select again to reset buffer).\n"
            "Confirm the trace only rises, then continue.",
        )

    def _arm_device(self):
        if not self._require_usb():
            return
        self.ui.armStatus.setText("Taking serial control and arming…")
        QtCore.QCoreApplication.processEvents()
        try:
            self.serial.take()
            self._send_arm_commands()
            self._apply_arm_result(self._verify_maxhold())
        except Exception as e:
            logging.exception("arm failed")
            QMessageBox.critical(
                self.ui,
                "Arm failed",
                f"{e}\n\nTip: click Stop on the main window first, then try Arm again.",
            )
            self.armed = False
            self.serial.release()
        self._update_nav()

    def _measure_sweep_time(self):
        if not self._require_usb():
            return
        self.ui.armStatus.setText("Measuring one sweep…")
        QtCore.QCoreApplication.processEvents()
        try:
            self.serial.take()
            self.serial.write("pause\r")
            self.serial.write("ultra on\r")
            self.serial.write(f"sweep start {START_FREQ_HZ}\r")
            self.serial.write(f"sweep stop {STOP_FREQ_HZ}\r")
            self.serial.write(f"rbw {RBW_KHZ}\r")
            self.serial.write("resume\r")
            t0 = time.time()
            time.sleep(0.5)
            self.serial.query("data 2\r")
            self.serial.query("data 2\r")
            elapsed = time.time() - t0
            hold_min = 3 if elapsed < 30 else 5
            self.ui.holdMinutes.setValue(hold_min)
            self.ui.armStatus.setText(
                f"Sweep probe ≈ {elapsed:.1f} s → hold set to {hold_min} minutes."
            )
            self.ui.armStatus.setStyleSheet("color: orange;")
        except Exception as e:
            logging.exception("sweep time probe failed")
            QMessageBox.critical(
                self.ui,
                "Sweep probe failed",
                f"{e}\n\nTip: click Stop on the main QtTinySA window, then retry.",
            )
            self.serial.release()
        else:
            # Keep FIFO stopped until Arm / Next / Cancel so the main sweep
            # cannot steal the port mid-setup.
            pass

    def _quiesce_app(self):
        self.serial.take()
        if self.usbCheck is not None and self.usbCheck.isActive():
            self.usbCheck.stop()
        # Close only the analyser under test; any other connected device keeps running
        self.serial.close_device()
        self._quiesced = True

    def _restore_app(self):
        if not self._quiesced:
            self.serial.release()
            return
        if self.usbCheck is not None:
            self.usbCheck.start(500)
        self.serial.reprobe()
        self.serial.restart_fifo()
        self._quiesced = False

    def _cleanup(self, restore=True):
        self.poll_timer.stop()
        self.countdown_timer.stop()
        if restore:
            self._restore_app()
        else:
            self.serial.release()
        self._save_prefs()

    # ------------------------------------------------------------------ unplug / hold / replug
    def _begin_unplug(self):
        self.hold_total_s = int(self.ui.holdMinutes.value()) * 60
        self.hold_remaining_s = self.hold_total_s
        self._quiesce_app()
        self.ui.unplugStatus.setText("Waiting for USB disconnect…")
        self.ui.unplugStatus.setStyleSheet("font-size: 12pt; color: orange;")
        self.ui.wizardStack.setCurrentIndex(PAGE_UNPLUG)
        self.poll_timer.start()

    def _on_poll_usb(self):
        idx = self.ui.wizardStack.currentIndex()
        ports = find_tinysa_ports()
        if idx == PAGE_UNPLUG:
            if not ports:
                self.poll_timer.stop()
                self.ui.unplugStatus.setText("USB disconnect detected.")
                self.ui.unplugStatus.setStyleSheet("font-size: 12pt; color: #0a0;")
                QtCore.QTimer.singleShot(800, self._begin_hold)
        elif idx == PAGE_REPLUG:
            if ports:
                self.poll_timer.stop()
                self.ui.replugStatus.setText(f"USB reconnect detected on {ports[0].device}.")
                self.ui.replugStatus.setStyleSheet("font-size: 12pt; color: #0a0;")
                QtCore.QTimer.singleShot(500, lambda: self._on_replugged(ports[0]))

    def _begin_hold(self):
        self.ui.wizardStack.setCurrentIndex(PAGE_HOLD)
        self.hold_remaining_s = self.hold_total_s
        self._update_countdown_ui()
        self.countdown_timer.start()

    def _update_countdown_ui(self):
        m, s = divmod(max(0, self.hold_remaining_s), 60)
        self.ui.countdownLabel.setText(f"{m:02d}:{s:02d}")
        done = self.hold_total_s - self.hold_remaining_s
        pct = int(100 * done / self.hold_total_s) if self.hold_total_s else 100
        self.ui.holdProgress.setValue(pct)

    def _on_countdown_tick(self):
        self.hold_remaining_s -= 1
        self._update_countdown_ui()
        if self.hold_remaining_s <= 0:
            self.countdown_timer.stop()
            self._begin_replug()

    def _abort_hold(self):
        self.countdown_timer.stop()
        reply = QMessageBox.question(
            self.ui,
            "Abort hold",
            "Abort the hold and proceed to replug / dump now?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if reply == QMessageBox.StandardButton.Yes:
            self._begin_replug()
        else:
            self.countdown_timer.start()

    def _begin_replug(self):
        self.ui.replugStatus.setText("Waiting for USB reconnect…")
        self.ui.replugStatus.setStyleSheet("font-size: 12pt; color: orange;")
        self.ui.wizardStack.setCurrentIndex(PAGE_REPLUG)
        self.poll_timer.start()

    def _on_replugged(self, port):
        self.serial.reprobe()
        if not self.serial.is_open():
            QMessageBox.critical(
                self.ui,
                "Reconnect failed",
                f"Analyser with serial {self.serial.sn} did not reappear on {port.device}.\n\n"
                "Check the cable and that the TinySA has powered up, then replug.",
            )
            self.poll_timer.start()
            return
        self.ui.wizardStack.setCurrentIndex(PAGE_DUMP)
        QtCore.QTimer.singleShot(300, self._dump_and_save)

    # ------------------------------------------------------------------ dump / save
    def _convert_trace(self, freq_hz, dbm):
        path = self.ui.probePath.currentIndex()
        gain = self.ui.ampGainDb.value()
        if path == 0:
            return compute_emc_tekbox(freq_hz, dbm, gain)
        return compute_emc_schwarzbeck(freq_hz, dbm, gain)

    def _convert_point(self, freq_hz: float, dbm: float):
        return self._convert_trace(freq_hz, dbm)

    def _read_trace(self):
        if not self.serial.is_open():
            raise RuntimeError("Serial port not open after reconnect")
        u = self.ui
        self.serial.write("pause\r")
        time.sleep(0.2)
        u.dumpStatus.setText("Reading frequencies…")
        QtCore.QCoreApplication.processEvents()
        freqs = parse_float_lines(self.serial.query("frequencies\r"))
        u.dumpStatus.setText("Reading data 2 (max hold)…")
        QtCore.QCoreApplication.processEvents()
        vals = parse_float_lines(self.serial.query("data 2\r"))
        if not freqs or not vals:
            raise RuntimeError(f"Empty dump: {len(freqs)} freqs, {len(vals)} vals")
        n = min(len(freqs), len(vals))
        return freqs[:n], vals[:n]

    def _write_maxhold_csv(self, freqs, vals, actual, E, dbuv, label: str, safe: str) -> Path:
        u = self.ui
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = self.session_dir / f"TinySA_{safe}_{ts}_MaxHold.csv"
        freqs_arr = np.asarray(freqs, dtype=float)
        dbm_arr = np.asarray(vals, dtype=float)
        actual = np.asarray(actual, dtype=float)
        E = np.asarray(E, dtype=float)
        dbuv = np.asarray(dbuv, dtype=float)
        e_str = np.where(np.isnan(E), "nan", np.char.mod("%.4f", E))
        d_str = np.where(np.isnan(dbuv), "nan", np.char.mod("%.4f", dbuv))
        with open(out, "w", newline="", encoding="utf-8") as f:
            f.write("# TinySA FCC Near-Field MaxHold\n")
            f.write(f"# label={label}\n")
            f.write(f"# dut={u.dutName.text().strip()}\n")
            f.write(f"# start_Hz={START_FREQ_HZ}\n")
            f.write(f"# stop_Hz={STOP_FREQ_HZ}\n")
            f.write(f"# rbw_kHz={RBW_KHZ}\n")
            f.write(f"# hold_seconds={self.hold_total_s}\n")
            f.write(f"# distance_cm={u.distanceCm.value()}\n")
            f.write(f"# probe_path={u.probePath.currentText()}\n")
            f.write(f"# amp_or_af={u.ampGainDb.value()}\n")
            f.write(f"# notes={u.notes.toPlainText().strip()}\n")
            f.write(f"# timestamp={datetime.now().isoformat(timespec='seconds')}\n")
            f.write(f"# source=QtTinySA_FCC_wizard\n")
            writer = csv.writer(f)
            writer.writerow(
                [
                    "Frequency_Hz",
                    "MaxHold_dBm",
                    "Actual_Signal_Level_dBm",
                    "E_uV_per_m",
                    "dBuV_per_m",
                ]
            )
            writer.writerows(
                zip(
                    np.char.mod("%.0f", freqs_arr),
                    np.char.mod("%.2f", dbm_arr),
                    np.char.mod("%.2f", actual),
                    e_str,
                    d_str,
                )
            )
        return out

    def _register_run(self, out: Path, label: str, freqs, vals, dbuv):
        u = self.ui
        run = {
            "label": label,
            "is_ambient": label.lower().startswith("ambient") or u.ambientRun.isChecked(),
            "path": str(out),
            "freq_hz": np.asarray(freqs, dtype=float),
            "dBuV": np.asarray(dbuv, dtype=float),
            "dbm": np.asarray(vals, dtype=float),
        }
        self.session_runs.append(run)
        xlsx_path = self._write_session_excel()
        n = len(run["freq_hz"])
        u.dumpStatus.setText(
            f"Saved {n} points → {out.name}"
            + (f" + {xlsx_path.name}" if xlsx_path else "")
        )
        self._clear_maxhold_buffer()
        self._restore_app()
        self._save_prefs()
        self._refresh_run_list()
        self._redraw_plot()
        QtCore.QTimer.singleShot(600, lambda: self.ui.wizardStack.setCurrentIndex(PAGE_RESULTS))

    def _dump_and_save(self):
        u = self.ui
        u.dumpStatus.setText("Sending pause…")
        QtCore.QCoreApplication.processEvents()
        try:
            freqs, vals = self._read_trace()
            label = u.runLabel.text().strip()
            if u.ambientRun.isChecked() and not label.lower().startswith("ambient"):
                label = "Ambient"
            actual, E, dbuv = self._convert_trace(freqs, vals)
            out = self._write_maxhold_csv(
                freqs, vals, actual, E, dbuv, label, sanitize_label(label)
            )
            self._register_run(out, label, freqs, vals, dbuv)
        except Exception as e:
            logging.exception("dump failed")
            u.dumpStatus.setText(f"Dump failed: {e}")
            QMessageBox.critical(self.ui, "Dump failed", str(e))
            self._restore_app()

    # ------------------------------------------------------------------ results
    def _session_excel_path(self) -> Path | None:
        if not self.session_dir:
            return None
        return self.session_dir / "FCC_Session_Measurements.xlsx"

    def _unique_sheet_name(self, label: str, used_names: set) -> str:
        name = re.sub(r"[\\/*?:\[\]]", "_", label)[:31] or "Run"
        base = name
        i = 2
        while name in used_names:
            name = f"{base[:28]}_{i}"
            i += 1
        used_names.add(name)
        return name

    def _append_run_sheet(self, wb, run, distance, device, used_names: set):
        name = self._unique_sheet_name(run["label"], used_names)
        ws = wb.create_sheet(title=name)
        ws.append(
            [
                "frequencies",
                "data2",
                "Power (dBm)",
                "Actual_Signal_Level_dBm",
                "E (uV/m)",
                "dBuV_per_m",
                "distance_cm",
                "device",
            ]
        )
        actual, E, _ = self._convert_trace(run["freq_hz"], run["dbm"])
        actual = np.asarray(actual, dtype=float)
        E = np.asarray(E, dtype=float)
        for fh, dbm, act, e, dbuv in zip(
            run["freq_hz"], run["dbm"], actual, E, run["dBuV"]
        ):
            ws.append(
                [
                    float(fh),
                    float(dbm),
                    float(dbm),
                    float(act),
                    float(e) if not np.isnan(e) else None,
                    float(dbuv) if not np.isnan(dbuv) else None,
                    distance,
                    device,
                ]
            )

    def _write_setup_sheet(self, wb):
        meta = wb.create_sheet(title="_Setup", index=0)
        meta.append(["Parameter", "Value"])
        meta.append(["span_start_Hz", START_FREQ_HZ])
        meta.append(["span_stop_Hz", STOP_FREQ_HZ])
        meta.append(["rbw_kHz", RBW_KHZ])
        meta.append(["rbw_note", "30 kHz RBW for finer harmonic resolution"])
        meta.append(["ultra_mode", "ultra on (PIN 4321 is on-device UI unlock only)"])
        meta.append(["probe_path", self.ui.probePath.currentText()])
        meta.append(["amp_or_af", self.ui.ampGainDb.value()])
        meta.append(["updated", datetime.now().isoformat(timespec="seconds")])

    def _write_session_excel(self) -> Path | None:
        """One workbook for the session: DUT sheets first, Ambient last (per doc)."""
        path = self._session_excel_path()
        if not path or not self.session_runs:
            return None
        try:
            from openpyxl import Workbook
        except ImportError:
            logging.info("openpyxl not installed — skipping session Excel update")
            return None

        wb = Workbook()
        wb.remove(wb.active)
        used_names = set()
        distance = self.ui.distanceCm.value()
        ordered = [r for r in self.session_runs if not r["is_ambient"]] + [
            r for r in self.session_runs if r["is_ambient"]
        ]
        for run in ordered:
            device = self.ui.dutName.text().strip() or run["label"]
            self._append_run_sheet(wb, run, distance, device, used_names)
        self._write_setup_sheet(wb)
        wb.save(path)
        return path

    def _display_label(self, run: dict, index: int) -> str:
        """Unique list/legend label when several runs share the same name."""
        same = [i for i, r in enumerate(self.session_runs) if r["label"] == run["label"]]
        if len(same) <= 1:
            return run["label"]
        return f"{run['label']} ({same.index(index) + 1})"

    def _refresh_run_list(self):
        self.ui.runList.blockSignals(True)
        self.ui.runList.clear()
        for i, run in enumerate(self.session_runs):
            item = QListWidgetItem(self._display_label(run, i))
            item.setFlags(item.flags() | QtCore.Qt.ItemFlag.ItemIsUserCheckable)
            item.setCheckState(QtCore.Qt.CheckState.Checked)
            item.setData(QtCore.Qt.ItemDataRole.UserRole, run["path"])
            self.ui.runList.addItem(item)
        self.ui.runList.blockSignals(False)

    def _on_run_visibility(self, _item):
        self._redraw_plot()

    def _visible_runs(self):
        visible = set()
        for i in range(self.ui.runList.count()):
            item = self.ui.runList.item(i)
            if item.checkState() == QtCore.Qt.CheckState.Checked:
                visible.add(item.data(QtCore.Qt.ItemDataRole.UserRole))
        return visible

    def _clear_plot_overlays(self, pw):
        for curve in list(self._plot_curves.values()):
            pw.removeItem(curve)
        self._plot_curves.clear()
        for item in list(pw.items()):
            if isinstance(item, pyqtgraph.InfiniteLine) and getattr(item, "_fcc_harmonic", False):
                pw.removeItem(item)
        legend = pw.plotItem.legend
        if legend is not None:
            legend.clear()
            legend.addItem(self._fcc_limit_curve, "FCC Class B Limit")

    def _draw_harmonic_lines(self, pw):
        carrier = self.ui.carrierMhz.value()
        n = 0
        while True:
            f_mhz = (2 * n + 1) * carrier
            if f_mhz > 1000:
                break
            if f_mhz >= 15:
                line = pyqtgraph.InfiniteLine(
                    pos=f_mhz,
                    angle=90,
                    pen=pyqtgraph.mkPen((200, 200, 200), width=0.5),
                )
                line._fcc_harmonic = True
                pw.addItem(line)
            n += 1

    def _plot_visible_runs(self, pw):
        visible = self._visible_runs()
        dut_colors = ["#1f77b4", "#ff7f0e", "#d62728", "#9467bd", "#8c564b", "#e377c2"]
        ambient_color = "#2ca02c"
        y_min, y_max = 10.0, 90.0
        plotted_any = False
        ordered = [
            (i, run)
            for i, run in enumerate(self.session_runs)
            if run["path"] in visible
        ]
        ordered.sort(key=lambda ir: (ir[1]["is_ambient"], ir[0]))
        dut_i = 0
        for i, run in ordered:
            freq_mhz = np.asarray(run["freq_hz"], dtype=float) / 1e6
            dBuV = np.asarray(run["dBuV"], dtype=float)
            label = self._display_label(run, i)
            if run["is_ambient"]:
                pen = pyqtgraph.mkPen(ambient_color, width=1.8)
            else:
                pen = pyqtgraph.mkPen(dut_colors[dut_i % len(dut_colors)], width=1.5)
                dut_i += 1
            curve = pw.plot(freq_mhz, dBuV, pen=pen, name=label)
            self._plot_curves[run["path"]] = curve
            finite = dBuV[np.isfinite(dBuV)]
            if finite.size:
                plotted_any = True
                y_min = min(y_min, float(np.nanmin(finite)))
                y_max = max(y_max, float(np.nanmax(finite)))
        return y_min, y_max, plotted_any

    def _redraw_plot(self):
        pw = self.ui.fccPlot
        self._clear_plot_overlays(pw)
        self._draw_harmonic_lines(pw)
        y_min, y_max, plotted_any = self._plot_visible_runs(pw)
        self._update_fcc_limit_curve()
        pw.setXRange(0, 1000, padding=0)
        if plotted_any:
            pad = max(5.0, 0.08 * (y_max - y_min))
            pw.setYRange(max(0.0, y_min - pad), y_max + pad, padding=0)
        else:
            pw.setYRange(10, 90, padding=0)

    def _new_run(self):
        # Ensure previous max-hold cannot contaminate the next arm
        try:
            if self._require_usb():
                self.serial.take()
                self.serial.write("pause\r")
                self.serial.write("calc off\r")
        except Exception as e:
            logging.info(f"New Run buffer clear: {e}")
        finally:
            self.serial.release()
        self.armed = False
        self.ui.runLabel.clear()
        self.ui.dutName.clear()
        self.ui.notes.clear()
        self.ui.ambientRun.setChecked(False)
        for chk in (
            self.ui.chkPhoto,
            self.ui.chkDutCharged,
            self.ui.chkTinyCharged,
            self.ui.chkPlacement,
        ):
            chk.setChecked(False)
        self.ui.wizardStack.setCurrentIndex(PAGE_CHECKLIST)
        self._update_nav()

    def _save_png(self):
        if not self.session_dir:
            return
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = self.session_dir / f"Glass_FCC_Plot_{ts}.png"
        exporter = pyqtgraph.exporters.ImageExporter(self.ui.fccPlot.plotItem)
        exporter.parameters()["width"] = 1600
        exporter.export(str(path))
        QMessageBox.information(self.ui, "Saved", f"PNG saved:\n{path}")

    def _export_excel(self):
        if not self.session_runs:
            QMessageBox.information(self.ui, "Export", "No runs to export.")
            return
        path = self._write_session_excel()
        if path is None:
            QMessageBox.critical(
                self.ui,
                "openpyxl missing",
                "Install openpyxl to export Excel:\n  pip install openpyxl",
            )
            return
        QMessageBox.information(
            self.ui,
            "Export",
            f"Session Excel updated (DUT sheets first, Ambient last):\n{path}",
        )
